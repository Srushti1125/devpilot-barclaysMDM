from app.retrieval.tree_index import (
    Block, blocks_from_markdown, build_tree, subtree_text,
)

DOC = "doc-1"


def _by_title(nodes):
    return {n.title: n for n in nodes}


def test_hierarchy_follows_heading_levels():
    blocks = blocks_from_markdown(
        "# Authentication\n"
        "## Password Reset\nReset via OTP.\n"
        "## MFA\nAdmins need a second factor.\n"
        "# Payments\n"
        "## Limits\nDaily cap.\n"
    )
    nodes = build_tree(DOC, blocks)
    t = _by_title(nodes)
    assert t["Password Reset"].parent_id == t["Authentication"].node_id
    assert t["MFA"].parent_id == t["Authentication"].node_id
    assert t["Limits"].parent_id == t["Payments"].node_id
    assert t["Payments"].parent_id is None
    assert t["Password Reset"].text == "Reset via OTP."


def test_skipped_level_attaches_to_nearest_open_ancestor():
    blocks = blocks_from_markdown("# A\n### Deep\nbody\n")
    t = _by_title(build_tree(DOC, blocks))
    assert t["Deep"].parent_id == t["A"].node_id


def test_parent_page_range_covers_descendants():
    blocks = [
        Block("Intro", page=1, heading_level=1, title="Intro"),
        Block("Child", page=2, heading_level=2, title="Child"),
        Block("body on p3", page=3),
        Block("Next", page=4, heading_level=1, title="Next"),
    ]
    t = _by_title(build_tree(DOC, blocks))
    assert (t["Intro"].page_start, t["Intro"].page_end) == (1, 3)
    assert (t["Child"].page_start, t["Child"].page_end) == (2, 3)
    assert (t["Next"].page_start, t["Next"].page_end) == (4, 4)


def test_preamble_text_is_kept_as_root_node():
    blocks = [Block("Front matter line", page=1)] + blocks_from_markdown("# Real\nbody")
    nodes = build_tree(DOC, blocks)
    assert nodes[0].title == "(preamble)"
    assert nodes[0].text == "Front matter line"
    assert nodes[0].parent_id is None
    real = _by_title(nodes)["Real"]
    assert real.parent_id is None
    assert real.ordinal != nodes[0].ordinal


def test_node_ids_are_deterministic_and_doc_scoped():
    blocks = blocks_from_markdown("# A\nx\n## B\ny\n")
    first = [n.node_id for n in build_tree("doc-1", blocks)]
    again = [n.node_id for n in build_tree("doc-1", blocks)]
    other = [n.node_id for n in build_tree("doc-2", blocks)]
    assert first == again
    assert set(first).isdisjoint(other)


def test_flat_document_becomes_low_structure_pseudo_sections():
    blocks = [Block("x" * 400, page=p) for p in range(1, 9)]
    nodes = build_tree(DOC, blocks, pseudo_section_chars=1000)
    assert len(nodes) > 1
    assert all(n.low_structure for n in nodes)
    assert all(n.parent_id is None for n in nodes)
    assert nodes[0].page_start == 1
    assert nodes[-1].page_end == 8


def test_subtree_text_includes_children_in_order():
    blocks = blocks_from_markdown("# A\nintro\n## B\nbee\n## C\nsee\n")
    a = _by_title(build_tree(DOC, blocks))["A"]
    text = subtree_text(a)
    assert text.index("intro") < text.index("bee") < text.index("see")


def test_docx_heading_styles_build_hierarchy(tmp_path):
    from docx import Document as DocxDocument
    from app.retrieval.tree_index import blocks_from_docx

    d = DocxDocument()
    d.add_heading("1 Authentication", level=1)
    d.add_heading("1.1 Password Reset", level=2)
    d.add_paragraph("Reset via OTP within 10 minutes.")
    d.add_heading("2 Payments", level=1)
    d.add_paragraph("Limits apply.")
    path = tmp_path / "reqs.docx"
    d.save(str(path))

    t = _by_title(build_tree(DOC, blocks_from_docx(str(path))))
    assert t["1.1 Password Reset"].parent_id == t["1 Authentication"].node_id
    assert t["1.1 Password Reset"].text == "Reset via OTP within 10 minutes."
    assert t["2 Payments"].text == "Limits apply."


def test_header_level_from_numbering_overrides_docling_level():
    from app.ingestion.docling_parser import _header_level
    assert _header_level("1 Authentication", 1) == 1
    assert _header_level("1.1 Password Reset", 1) == 2
    assert _header_level("2.2.3 Deep", 1) == 3
    assert _header_level("Unnumbered Title", 2) == 2
    assert _header_level("Unnumbered Title", None) == 1
    assert _header_level("7.1.1.1.1.1.1 Too deep", 1) == 6
